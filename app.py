import io
import os
import re
import tempfile
from decimal import Decimal
from datetime import datetime

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, numbers
from openpyxl.utils import get_column_letter
import xlsxwriter

import db
from queries import get_salary_by_cert, get_cross_unit_salary_rows

app = Flask(__name__)

CERT_RE = re.compile(r"^[0-9]{17}[0-9X]$")

# Money fields to format as 2-decimal floats
MONEY_FIELDS = {
    "ATC93X", "ATC933", "ATC934", "ATC935", "ATC936", "ATC93K",
    "ATC93AA", "ATC93D", "ATC93C", "ATC93AB",
}


def _serialize_row(row):
    """Convert Oracle row dict to JSON-safe dict."""
    out = {}
    for k, v in row.items():
        if isinstance(v, Decimal):
            out[k] = float(v)
        elif v is None:
            out[k] = None
        else:
            out[k] = v
    return out


def _validate_cert(cert_raw):
    cert = (cert_raw or "").strip().upper()
    if not CERT_RE.match(cert):
        return None, "身份证号格式错误，应为18位（末位可为X）"
    return cert, None


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/salary")
def api_salary():
    cert_raw = request.args.get("cert", "")
    cert, err = _validate_cert(cert_raw)
    if err:
        return jsonify({"ok": False, "error": err}), 400

    conn = None
    try:
        conn = db.get_connection()
        rows = get_salary_by_cert(conn, cert)
    except Exception as e:
        return jsonify({"ok": False, "error": f"数据库查询失败: {e}"}), 500
    finally:
        db.release_connection(conn)

    if not rows:
        return jsonify({"ok": True, "cert": cert, "name": None, "count": 0, "rows": []})

    name = rows[0].get("AAC003")
    aac001 = rows[0].get("AAC001")
    serialized = []
    for r in rows:
        s = _serialize_row(r)
        s["status_label"] = r["STATUS_LABEL"]
        s["status_raw_min"] = r["STATUS_RAW_MIN"]
        s["status_raw_max"] = r["STATUS_RAW_MAX"]
        s["jingban_month"] = r["JINGBAN_MONTH"]
        serialized.append(s)

    return jsonify({
        "ok": True,
        "cert": cert,
        "name": name,
        "aac001": aac001,
        "count": len(serialized),
        "rows": serialized,
    })


# Excel export column spec: (header, field_key, is_money)
EXCEL_COLUMNS = [
    ("姓名", "AAC003", False),
    ("身份证号", "AAC002", False),
    ("个人编号", "AAC001", False),
    ("员工号", "AAE00B", False),
    ("结算单元", "ATB931", False),
    ("工资所属年月", "ATC931", False),
    ("发放年月(ATC932)", "ATC932", False),
    ("发放批次", "ATC937", False),
    ("本次基本工资(ATC93X)", "ATC93X", True),
    ("本次应发工资(ATC933)", "ATC933", True),
    ("本次扣款(ATC934)", "ATC934", True),
    ("本次补发(ATC935)", "ATC935", True),
    ("本次免税(ATC936)", "ATC936", True),
    ("本次年终奖(ATC93K)", "ATC93K", True),
    ("本次工资总额(ATC93AA)", "ATC93AA", True),
    ("本次个人所得税(ATC93D)", "ATC93D", True),
    ("本次实发金额(ATC93C)", "ATC93C", True),
    ("累计应税金额(ATC93AB)", "ATC93AB", True),
    ("发放状态", "status_label", False),
    ("状态原值(ATC8M3_MIN)", "status_raw_min", False),
    ("经办年月(ATC8G7)", "jingban_month", False),
]


@app.route("/api/salary/export")
def api_salary_export():
    cert_raw = request.args.get("cert", "")
    cert, err = _validate_cert(cert_raw)
    if err:
        return jsonify({"ok": False, "error": err}), 400

    conn = None
    try:
        conn = db.get_connection()
        rows = get_salary_by_cert(conn, cert)
    except Exception as e:
        return jsonify({"ok": False, "error": f"数据库查询失败: {e}"}), 500
    finally:
        db.release_connection(conn)

    wb = Workbook()
    ws = wb.active
    ws.title = "工资明细"

    # Header row
    bold = Font(bold=True)
    for col_idx, (header, _, _) in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold

    # Data rows
    for row_idx, r in enumerate(rows, start=2):
        for col_idx, (_, field, is_money) in enumerate(EXCEL_COLUMNS, start=1):
            val = r.get(field)
            if isinstance(val, Decimal):
                val = float(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if is_money and val is not None:
                cell.number_format = "0.00"

    # Freeze top row
    ws.freeze_panes = "A2"

    # Auto-width (approximate)
    for col_idx, (header, _, _) in enumerate(EXCEL_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(header) * 1.8, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"工资明细_{cert}.xlsx",
    )


MONTH_RE = re.compile(r"^\d{4}(0[1-9]|1[0-2])$")
VALID_STATUSES = {"unpaid", "not_paid", "all"}
VALID_SCOPES = {"cross_unit", "all"}
EXCEL_MAX_DATA_ROWS = 1_048_575  # 1,048,576 total rows minus 1 header
STATUS_LABELS_EXPORT = {
    "unpaid": "仅未发放",
    "not_paid": "未发工资",
    "all": "全部状态",
}
SCOPE_FILENAME_PREFIX = {
    "cross_unit": "跨结算单元工资明细",
    "all": "全部人员工资明细",
}


@app.route("/api/cross-unit/export")
def api_cross_unit_export():
    from schema_cn import EXPORT_COLUMNS

    month_raw = request.args.get("month", "").strip()
    pay_month_raw = request.args.get("pay_month", "").strip()
    status = request.args.get("status", "unpaid").strip()
    scope = request.args.get("scope", "cross_unit").strip()

    if month_raw:
        if not MONTH_RE.match(month_raw):
            return jsonify({"ok": False, "error": "发薪年月格式错误，应为YYYYMM（如202609）"}), 400
        month = int(month_raw)
    else:
        month = None

    if pay_month_raw:
        if not MONTH_RE.match(pay_month_raw):
            return jsonify({"ok": False, "error": "发放年月格式错误，应为YYYYMM（如202510）"}), 400
        pay_month = int(pay_month_raw)
    else:
        pay_month = None

    if status not in VALID_STATUSES:
        return jsonify({"ok": False, "error": "状态筛选参数错误"}), 400

    if scope not in VALID_SCOPES:
        return jsonify({"ok": False, "error": "查询范围参数错误"}), 400

    # Fast path: reject unfiltered scope=all before the heavy query
    if scope == "all" and month is None and pay_month is None and status == "all":
        conn = None
        try:
            conn = db.get_connection()
            cursor = conn.cursor()
            try:
                cursor.execute("SELECT COUNT(*) FROM TC93")
                total = cursor.fetchone()[0]
            finally:
                cursor.close()
        except Exception as e:
            return jsonify({"ok": False, "error": f"数据库查询失败: {e}"}), 500
        finally:
            db.release_connection(conn)
        if total > EXCEL_MAX_DATA_ROWS:
            return jsonify({"ok": False, "error": "数据量超过 Excel 单表上限（1048576 行），请增加筛选条件"}), 400

    conn = None
    tmp_path = None
    try:
        conn = db.get_connection()
        columns, row_iter = get_cross_unit_salary_rows(
            conn, month=month, pay_month=pay_month, status=status, scope=scope
        )

        tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(tmp_fd)

        wb = xlsxwriter.Workbook(tmp_path, {"constant_memory": True})
        ws = wb.add_worksheet("跨结算单元工资明细")

        bold = wb.add_format({"bold": True})
        money_fmt = wb.add_format({"num_format": "0.00"})
        date_fmt = wb.add_format({"num_format": "yyyy-mm-dd"})

        for col_idx, col_spec in enumerate(EXPORT_COLUMNS):
            ws.write(0, col_idx, col_spec["header"], bold)

        ws.freeze_panes("A2")
        ws.set_column(0, len(EXPORT_COLUMNS) - 1, 14)

        wide_cols = {"AAC002", "BBB001", "ATB931", "AAB004", "AAB004_1"}
        for col_idx, col_spec in enumerate(EXPORT_COLUMNS):
            if col_spec["field"] in wide_cols:
                ws.set_column(col_idx, col_idx, 28)

        row_idx = 1
        overflow = False
        for row_data in row_iter:
            if row_idx > EXCEL_MAX_DATA_ROWS:
                overflow = True
                break
            for col_idx, (value, col_spec) in enumerate(zip(row_data, EXPORT_COLUMNS)):
                if value is None:
                    continue
                kind = col_spec["kind"]
                if kind == "str":
                    ws.write_string(row_idx, col_idx, str(value))
                elif kind == "id":
                    ws.write_string(row_idx, col_idx, str(int(value)))
                elif kind == "money":
                    ws.write_number(row_idx, col_idx, float(value), money_fmt)
                elif kind == "int":
                    ws.write_number(row_idx, col_idx, int(value))
                elif kind == "date":
                    ws.write_datetime(row_idx, col_idx, value, date_fmt)
            row_idx += 1

        wb.close()

        if overflow:
            # Fail closed: delete temp file and return 400
            try:
                if tmp_path and os.path.exists(tmp_path):
                    os.remove(tmp_path)
            except OSError:
                pass
            return jsonify({"ok": False, "error": "数据量超过 Excel 单表上限（1048576 行），请增加筛选条件"}), 400

        month_str = str(month) if month else "全部"
        pay_month_str = str(pay_month) if pay_month else "全部"
        status_str = STATUS_LABELS_EXPORT[status]
        prefix = SCOPE_FILENAME_PREFIX[scope]
        filename = f"{prefix}_发薪{month_str}_发放{pay_month_str}_{status_str}.xlsx"

        resp = send_file(
            tmp_path,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

        def cleanup():
            try:
                if tmp_path and os.path.exists(tmp_path):
                    os.remove(tmp_path)
            except OSError:
                pass

        resp.call_on_close(cleanup)
        return resp

    except Exception as e:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        return jsonify({"ok": False, "error": f"数据库查询失败: {e}"}), 500
    finally:
        db.release_connection(conn)


if __name__ == "__main__":
    db.init_db()
    app.run(host="0.0.0.0", port=5001, debug=False, threaded=True)
